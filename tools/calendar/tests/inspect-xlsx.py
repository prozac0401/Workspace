"""Read-only independent calendar + ZIP/OOXML/openpyxl verification. Never authors XLSX."""
from pathlib import Path
import calendar
import datetime as dt
import json
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from openpyxl import load_workbook

path = (Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[3] / 'artifacts/calendar/calendar_2027.xlsx').resolve()
model_path = Path(sys.argv[2]) if len(sys.argv) > 2 else path.with_name('model_' + path.stem.split('_')[-1] + '.json')
model = json.loads(model_path.read_text(encoding='utf-8'))
year = model['config']['year']
assert zipfile.is_zipfile(path), 'Real ZIP-based XLSX required; HTML/CSV is forbidden'
wb = load_workbook(path)
assert wb.sheetnames == [f'{year}년 달력'], wb.sheetnames
ws = wb.active
assert ws.max_column == 7
assert not ws.protection.sheet
assert ws.sheet_view.showGridLines is False
assert ws.print_options.gridLines is False
assert ws.page_setup.orientation == 'portrait'
assert str(ws.page_setup.paperSize) == '9'
assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0
assert ws.sheet_properties.pageSetUpPr.fitToPage is True
assert ws.print_title_rows is None
assert ws.column_dimensions['A'].min == 1 and ws.column_dimensions['A'].max == 7
assert ws.column_dimensions['A'].width == 18
assert ws['A1'].value == f'{year}년 달력'
assert ws['A2'].value == model['metadataText']
assert ws['A3'].value == model['legendText']
assert '모의 테스트' in ws['A2'].value
assert '자동 갱신되지 않습니다' in ws['A3'].value

ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
colors = {'black': 'FF17202A', 'red': 'FF9C2436', 'blue': 'FF2158A6', 'purple': 'FF70308C'}
dates = {}
for month in model['months']:
    for day in month['slots']:
        if day['date']:
            assert day['date'] not in dates, 'Model repeated a date'
            dates[day['date']] = day

with zipfile.ZipFile(path) as z:
    assert z.testzip() is None, 'ZIP CRC failure'
    names = z.namelist()
    assert len(names) == len(set(names)) == 6
    assert names.count('xl/worksheets/sheet1.xml') == 1
    assert not any(term in name.lower() for name in names for term in ['vba', 'externallink', 'connection'])
    for name in names:
        xml = ET.fromstring(z.read(name))
        assert not any(node.tag.split('}')[-1] in ['f', 'sheetProtection', 'externalReferences', 'connections'] for node in xml.iter())
        assert not any(node.attrib.get('TargetMode') == 'External' for node in xml.iter())
    sheet = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    cells = {c.attrib['r']: c for c in sheet.findall('.//s:c', ns)}
    merges = ['A1:G1', 'A2:G2', 'A3:G3']
    start = 5
    blanks = 0
    date_count = 0
    expected_breaks = []
    for month in range(1, 13):
        offset_mon, days = calendar.monthrange(year, month)
        offset = (offset_mon + 1) % 7
        weeks = (offset + days + 6) // 7
        assert model['months'][month - 1]['weeks'] == weeks
        assert model['months'][month - 1]['startRow'] == start
        assert ws.cell(start, 1).value == f'{year}년 {month}월'
        assert [ws.cell(start + 1, c).value for c in range(1, 8)] == list('일월화수목금토')
        merges.append(f'A{start}:G{start}')
        if month > 1:
            expected_breaks.append(start - 1)
        for k in range(weeks * 7):
            number = k - offset + 1
            row, col = start + 2 + 4 * (k // 7), k % 7 + 1
            cell = ws.cell(row, col)
            if 1 <= number <= days:
                iso = dt.date(year, month, number).isoformat()
                day = dates[iso]
                assert day['excelRow'] == row and day['excelColumn'] == col
                # openpyxl preserves the escaped underscore used by Excel's _xNNNN_ syntax.
                expected = re.sub(r'(_x[0-9a-fA-F]{4}_)', lambda m: '_x005F_' + m[0][1:], day['displayText'])
                assert cell.value == expected, (iso, cell.coordinate, cell.value, expected)
                assert cell.data_type == 's', (iso, cell.data_type)
                assert cell.font.color.rgb == colors[day['styleToken']]
                assert cell.alignment.wrapText and cell.alignment.vertical == 'top'
                assert 34 <= ws.row_dimensions[row].height <= 409
                date_count += 1
            else:
                assert cell.value is None, 'Adjacent-month date must stay empty'
                assert len(cells[cell.coordinate]) == 0
            for delta in range(1, 4):
                blank = ws.cell(row + delta, col)
                assert blank.value is None and blank.data_type != 'f'
                assert blank.coordinate not in ws.merged_cells
                assert ws.row_dimensions[blank.row].height == 22
                element = cells[blank.coordinate]
                assert len(element) == 0 and 't' not in element.attrib, 'Writing cell contains value/formula/string/space'
                assert element.attrib.get('s') is not None, 'Styled editable blank expected'
                if 1 <= number <= days:
                    blanks += 1
        end = start + 1 + 4 * weeks
        start += 3 + 4 * weeks
    assert sorted(str(x) for x in ws.merged_cells.ranges) == sorted(merges)
    assert [b.id for b in ws.row_breaks.brk] == expected_breaks
    assert len(expected_breaks) == 11
    assert ws.max_row == end
    assert f'$A$1:$G${end}' in str(ws.print_area)
    assert date_count == len(dates) == (366 if calendar.isleap(year) else 365)
    assert blanks == date_count * 3
    for row in ws:
        for cell in row:
            assert cell.data_type != 'f'
    for prefix in ['=SUM(A1:A2)', '+1', '-1', '@test']:
        assert any(prefix in str(c.value) and c.data_type == 's' for row in ws for c in row)

report = {'status': 'PASS', 'engine': 'openpyxl 3.1.5 + independent Python calendar + ZIP/XML (not Excel)',
          'year': year, 'dates': date_count, 'independentBlankCells': blanks, 'sheets': 1, 'months': 12,
          'columns': 7, 'titleMerges': 15, 'manualPageBreaks': 11, 'formulas': 0, 'datasetKind': 'fixture'}
path.with_name(f'inspection_{year}.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps(report, ensure_ascii=False))
